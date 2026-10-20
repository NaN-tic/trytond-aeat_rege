from decimal import Decimal
from io import BytesIO

from trytond.i18n import gettext
from trytond.pool import Pool, PoolMeta
from trytond.transaction import Transaction


class TaxesByInvoice(metaclass=PoolMeta):
    __name__ = 'account_reports.taxes_by_invoice'

    @classmethod
    def prepare(cls, data, checker=None):
        from trytond.modules.html_report.engine import DualRecord

        records, parameters = super().prepare(data, checker)
        pool = Pool()
        Invoice = pool.get('account.invoice')
        InvoiceLine = pool.get('account.invoice.line')
        InvoiceTax = pool.get('account.invoice.tax')
        Currency = pool.get('currency.currency')
        Period = pool.get('account.period')
        Tax = pool.get('account.tax')

        domain = [('type', '=', 'line'), ('invoice.move', '!=', None),
            ('invoice.type', '=', 'out' if data['partner_type'] == 'customers' else 'in')]
        if data.get('company'):
            domain.append(('invoice.company', '=', data['company']))
        for name, operator in [('start_date', '>='), ('end_date', '<=')]:
            if data.get(name):
                domain.append(('invoice.move.date', operator, data[name]))
        if not data.get('start_date') and not data.get('end_date'):
            periods = data.get('periods') or [p.id for p in Period.search([
                        ('fiscalyear', '=', data['fiscalyear'])])]
            domain.append(('invoice.move.period', 'in', periods))
        for name, operator in [('parties', 'in'), ('excluded_parties', 'not in')]:
            if data.get(name):
                domain.append(('invoice.party', operator, data[name]))
        if data.get('taxes'):
            domain.append(('taxes', 'in', data['taxes']))
        if data['tax_type'] == 'invoiced':
            domain.append(['OR',
                    [('quantity', '>=', 0), ('unit_price', '>=', 0)],
                    [('quantity', '<=', 0), ('unit_price', '<=', 0)]])
        elif data['tax_type'] == 'refunded':
            domain.append(['OR',
                    [('quantity', '>', 0), ('unit_price', '<', 0)],
                    [('quantity', '<', 0), ('unit_price', '>', 0)]])
        lines = InvoiceLine.search(domain, order=[('invoice', 'ASC'), ('id', 'ASC')])
        invoices = {line.invoice.id: line.invoice for line in lines}
        invoices = {id_: invoice for id_, invoice in invoices.items()
            if invoice._get_rege(advanced=True)}
        if not invoices:
            return records, parameters

        # Replace only automatic VAT and the base report's synthetic
        # non-deductible rows. Manual adjustments and other taxes are retained.
        for key, rows in list(records.items()):
            rows = [row for row in rows
                if row.invoice.raw.id not in invoices or not (
                    (row.raw.id or -1) < 0 or (
                        not row.raw.manual and row.tax
                        and row.tax.raw.tax_kind == 'vat'))]
            if rows:
                records[key] = rows
            else:
                del records[key]
        line_taxes = {id_: cls._rege_line_taxes(invoice)
            for id_, invoice in invoices.items()}
        deductible_taxes = {id_: cls._rege_line_taxes(invoice, deductible=True)
            for id_, invoice in invoices.items()}
        fiscal_groups = {}
        group_id = min([0] + [row.tax.raw.id or 0 for rows in records.values()
                for row in rows if row.tax]) - 1
        row_id = min([0] + [row.raw.id or 0 for rows in records.values()
                for row in rows]) - 1
        details = {}
        for line in lines:
            if line.invoice.id not in invoices:
                continue
            invoice = line.invoice
            for tax_key, value in line_taxes[invoice.id].get(line.id, {}).items():
                tax = value.tax
                if (tax.tax_kind != 'vat'
                        or (data.get('taxes') and tax.id not in data['taxes'])):
                    continue
                base = invoice.currency.round(
                    Decimal(str(line.quantity or 0)) * (line.unit_price or Decimal(0)))
                with Transaction().set_context(date=invoice.currency_date):
                    company_base = Currency.compute(invoice.currency,
                        base, invoice.company.currency)
                    company_amount = Currency.compute(invoice.currency,
                        value.amount, invoice.company.currency)
                    cost_base = Currency.compute(invoice.currency,
                        value.base, invoice.company.currency)
                rate = (line.taxes_deductible_rate if invoice.type == 'in'
                    and line.taxes_deductible_rate is not None else Decimal(1))
                parts = []
                if not value.base and not value.amount and tax.rate:
                    parts.append(('zero', base, value.amount,
                            company_base, company_amount, cost_base))
                elif rate == 0 or rate == 1:
                    parts.append(('deductible' if rate else 'non_deductible',
                            base, value.amount, company_base, company_amount, cost_base))
                else:
                    deducted = deductible_taxes[invoice.id].get(line.id, {}).get(tax_key)
                    deducted_amount = deducted.amount if deducted else Decimal(0)
                    deducted_cost = deducted.base if deducted else Decimal(0)
                    deducted_base = invoice.currency.round(base * rate)
                    with Transaction().set_context(date=invoice.currency_date):
                        deducted_company_base = Currency.compute(invoice.currency,
                            deducted_base, invoice.company.currency)
                        deducted_company_amount = Currency.compute(invoice.currency,
                            deducted_amount, invoice.company.currency)
                        deducted_company_cost = Currency.compute(invoice.currency,
                            deducted_cost, invoice.company.currency)
                    parts.append(('deductible', deducted_base, deducted_amount,
                            deducted_company_base, deducted_company_amount,
                            deducted_company_cost))
                    # Use the remainder so rounding cannot duplicate or lose
                    # commercial bases, fiscal bases or VAT between groups.
                    parts.append(('non_deductible', base - deducted_base,
                            value.amount - deducted_amount,
                            company_base - deducted_company_base,
                            company_amount - deducted_company_amount,
                            cost_base - deducted_company_cost))
                for kind, part_base, amount, c_base, c_amount, c_cost in parts:
                    report_tax = tax
                    if kind != 'deductible':
                        group_key = (kind, tax.rate, invoice.company.id,
                            tax.id if tax.type != 'percentage' else None)
                        if group_key not in fiscal_groups:
                            percentage = (format((tax.rate * 100).normalize(), 'f')
                                if tax.rate is not None else None)
                            if kind == 'zero':
                                label = gettext('aeat_rege.msg_report_zero_cost_tax',
                                    rate=percentage)
                            else:
                                label = gettext('account_reports.msg_not_deductible_tax')
                                if percentage is not None:
                                    label += ' (%s%%)' % percentage
                            existing = next((key for key in records
                                    if isinstance(key, Tax) and (key.id or -1) < 0
                                    and key.name == label and key.rate == tax.rate
                                    and key.company == invoice.company), None)
                            if existing is not None:
                                fiscal_groups[group_key] = existing
                            else:
                                fiscal_groups[group_key] = Tax(group_id, name=label,
                                    rate=tax.rate, type=tax.type, company=invoice.company)
                                group_id -= 1
                        report_tax = fiscal_groups[group_key]
                    fiscal_line = InvoiceTax(row_id, invoice=invoice, tax=report_tax,
                        account=value.account, base=part_base, amount=amount,
                        currency=invoice.currency, company_base=c_base,
                        company_amount=c_amount, company_base_cache=c_base,
                        company_amount_cache=c_amount)
                    row_id -= 1
                    key = (invoice.move.period
                        if data['grouping'] == 'invoice' else report_tax)
                    records.setdefault(key, []).append(DualRecord(fiscal_line))
                    details[(invoice.id, fiscal_line.id, report_tax.id)] = {
                        'cost_base': c_cost, 'deduction_kind': kind,
                        'description': line.description or (
                            line.product.rec_name if line.product else ''),
                        }

        totals = dict.fromkeys(
            ['total_untaxed', 'total_tax', 'total', 'cost_base'], Decimal(0))
        tax_totals = {}
        for key, rows in records.items():
            rows.sort(key=lambda row: (
                row.invoice.raw.invoice_date or row.invoice.move.raw.date,
                row.invoice.raw.id))
            subtotal = dict.fromkeys(totals, Decimal(0))
            for row in rows:
                invoice = row.invoice.raw
                if (invoice.state == 'cancelled' and (not invoice.cancel_move
                        or not isinstance(invoice.cancel_move.origin, Invoice))):
                    continue
                subtotal['total_untaxed'] += row.raw.company_base
                subtotal['total_tax'] += row.raw.company_amount
                subtotal['total'] += row.raw.company_base + row.raw.company_amount
                detail = cls._rege_detail(row, details)
                subtotal['cost_base'] += detail.get('cost_base', Decimal(0))
            tax_totals[key] = subtotal
            for name, value in subtotal.items():
                totals[name] += value
        parameters.update(rege_details=details, totals=totals, tax_totals=tax_totals,
            rege_only=bool(details) and all(cls._rege_detail(row, details)
                for rows in records.values() for row in rows))
        return records, parameters

    @classmethod
    def _rege_line_taxes(cls, invoice, deductible=False):
        """Allocate fiscal taxes using the invoice's rounding policy."""
        with Transaction().set_context(_deductible_rate=None if deductible else 1):
            invoice_taxes = invoice._get_rege_taxes()
            line_taxes = {line.id: line._get_taxes()
                for line in invoice.lines if line.type == 'line'}
        for key, total in invoice_taxes.items():
            values = [taxes[key] for taxes in line_taxes.values() if key in taxes]
            if values:
                values[-1].amount += total.amount - sum(v.amount for v in values)
                values[-1].base += total.base - sum(v.base for v in values)
        return line_taxes

    @staticmethod
    def _rege_detail(line, details):
        return details.get((line.invoice.raw.id, line.raw.id,
                line.tax.raw.id if line.tax else None), {})

    @staticmethod
    def _base_label(parameters):
        return gettext('aeat_rege.msg_report_commercial_base'
            if parameters.get('rege_only') else 'aeat_rege.msg_report_mixed_base')

    @classmethod
    def _tax_label(cls, line, details):
        if not line.tax:
            return ' --- '
        tax = line.tax.raw
        detail = cls._rege_detail(line, details)
        if detail.get('deduction_kind') == 'zero':
            return gettext('aeat_rege.msg_report_zero_cost_tax',
                rate=format((tax.rate * 100).normalize(), 'f'))
        return tax.name

    @classmethod
    def _build_pdf(cls, records, parameters, margin, checker):
        from trytond.modules.account_reports.common import AccountReportsPDF
        from trytond.modules.html_report.engine import render
        from trytond.modules.xgettext import _

        details = parameters.get('rege_details')
        if not details:
            return super()._build_pdf(records, parameters, margin, checker)
        p = parameters
        title = _('Taxes By Invoice and Period') if p['grouping'] else _('Taxes By Invoice')
        periods = (_('Periods: %(periods)s') % p if p['periods'] else _('All Periods'))
        filters = ['%s: %s' % (p['company_vat_label'], p['company_vat']),
            _('Fiscal Year: %(fiscal_year)s %(periods)s') % {
                'fiscal_year': p['fiscal_year'], 'periods': periods}]
        if p['start_date']:
            filters += [_('Initial posting date: %(start_date)s') % p,
                _('Final posting date: %(end_date)s') % p]
        filters += [_('Parties: %(parties)s') % p if p['parties'] else _('All Parties'),
            _('Cancelled invoices are shown in grey. Invoices '
                'without a cancelled move or a cancelled move not '
                'related to an invoice are not added to the total.')]
        pdf = AccountReportsPDF(title, p['company'], filters,
            [_('Move Date'), _('Account'), '', _('NIF'), _('Invoice'), _('Date'),
                cls._base_label(p), '', _('Tax'), _('Base + Tax'), _('Total Invoice'),
                gettext('aeat_rege.msg_report_cost_column'),
                gettext('aeat_rege.msg_report_operation')],
            [7, 6, 14, 8, 9, 7, 9, 12, 8, 9, 10, 9, 14],
            ['L'] * 6 + ['R', 'L', 'R', 'R', 'R', 'R', 'L'], margin, checker)
        pdf.add_page()
        total_widths = [sum(pdf.widths[:6])] + pdf.widths[6:]
        for index, (key, lines) in enumerate(records.items()):
            if p['grouping'] and index:
                pdf.add_page()
            pdf.row([key.name if hasattr(key, 'name') else key.rec_name],
                widths=[pdf.epw], aligns=['L'], bold=True, fill=True)
            digits = key.company.currency.digits
            if not p['totals_only']:
                previous_invoice = None
                for line in lines:
                    invoice = line.invoice.raw
                    if invoice.id != previous_invoice:
                        tax_identifier = (invoice.party_tax_identifier
                            or invoice.party.tax_identifier)
                        values = [render(invoice.move.date), line.account.raw.code,
                            invoice.party.rec_name,
                            tax_identifier.code if tax_identifier else '',
                            ('*' if invoice.state == 'cancelled' else '') + (invoice.number or ''),
                            render(invoice.invoice_date)]
                        invoice_total = render(invoice.company_total_amount, digits=digits)
                    else:
                        values = [''] * 6
                        invoice_total = ''
                    detail = cls._rege_detail(line, details)
                    values += [render(line.raw.company_base, digits=digits),
                        cls._tax_label(line, details),
                        render(line.raw.company_amount, digits=digits),
                        render(line.raw.company_base + line.raw.company_amount, digits=digits),
                        invoice_total,
                        render(detail['cost_base'], digits=digits) if detail else '',
                        detail.get('description', '')]
                    pdf.row(values, fill=invoice.state == 'cancelled')
                    previous_invoice = invoice.id
            totals = p['tax_totals'].get(key)
            if totals:
                pdf.row([_('Total Period') if p['grouping'] else _('Total'),
                    render(totals['total_untaxed'], digits=digits), '',
                    render(totals['total_tax'], digits=digits),
                    render(totals['total'], digits=digits), '',
                    render(totals['cost_base'], digits=digits), ''],
                    widths=total_widths, aligns=['R'] * 8, bold=True, bottom=True)
            if p['grouping'] and index == len(records) - 1:
                totals = p['totals']
                pdf.row([_('Total'), render(totals['total_untaxed'], digits=digits), '',
                    render(totals['total_tax'], digits=digits),
                    render(totals['total'], digits=digits), '',
                    render(totals['cost_base'], digits=digits), ''],
                    widths=total_widths, aligns=['R'] * 8, bold=True, bottom=True)
        return bytes(pdf.output())


class TaxesByInvoiceXlsx(metaclass=PoolMeta):
    __name__ = 'account_reports.taxes_by_invoice_xlsx'

    @classmethod
    def get_content(cls, ids, data):
        # The base XLSX report calls its concrete Python class. Use the Pool
        # so that the optional REGE preparation also applies to spreadsheets.
        Report = Pool().get('account_reports.taxes_by_invoice', type='report')
        with Transaction().set_context(active_test=False):
            records, parameters = Report.prepare(data)
        if not records:
            parameters['records_found'] = False
            records['no_records'] = ''
        with Transaction().set_context(**cls._xlsx_context()):
            return cls._build_workbook(records, parameters)

    @classmethod
    def _build_workbook(cls, records, parameters):
        content = super()._build_workbook(records, parameters)
        details = parameters.get('rege_details')
        if not details:
            return content
        from openpyxl import load_workbook
        from trytond.modules.account_reports.xlsx import save_workbook
        from trytond.modules.xgettext import _

        Report = Pool().get('account_reports.taxes_by_invoice', type='report')
        workbook = load_workbook(BytesIO(content))
        sheet = workbook.active
        headers = [row[0].row for row in sheet.iter_rows()
            if row[0].value == _('Move Date') and row[10].value == _('Total Invoice')]
        groups = list(records.items())
        for group_index, (header, (key, lines)) in enumerate(zip(headers, groups)):
            digits = key.company.currency.digits
            sheet.cell(header, 7, Report._base_label(parameters))
            sheet.cell(header, 12, gettext('aeat_rege.msg_report_cost_column'))
            sheet.cell(header, 13, gettext('aeat_rege.msg_report_operation'))
            index = header + 2
            if not parameters['totals_only']:
                for line in lines:
                    detail = Report._rege_detail(line, details)
                    sheet.cell(index, 8, Report._tax_label(line, details))
                    if detail:
                        sheet.cell(index, 12, float(detail['cost_base']))
                        sheet.cell(index, 13, detail['description'])
                    index += 1
            if parameters['tax_totals'].get(key):
                sheet.cell(index, 12, float(parameters['tax_totals'][key]['cost_base']))
                index += 1
            if parameters['grouping'] and group_index == len(groups) - 1:
                sheet.cell(index, 12, float(parameters['totals']['cost_base']))
            for row in range(header + 2, index + 1):
                sheet.cell(row, 12).number_format = '0.' + '0' * digits
        return save_workbook(workbook)
